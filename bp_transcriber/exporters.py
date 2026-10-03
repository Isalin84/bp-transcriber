"""Экспорт транскрипта (dict по контракту §2.1) в txt, md, docx, srt, vtt, json.

Отображаемые имена спикеров берутся из ``transcript["speakers"]``.
"""

from __future__ import annotations

import json
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from collections.abc import Iterable

from .settings import FORMATS, atomic_write_text

NAVY = "0B1D3A"
STEEL_BLUE = "1E3A5F"
GRAY = "6B7280"
FOOTER_TEXT = "Создано в BP Transcriber · bestpracticeai.ru"


class ExportError(Exception):
    """Ошибка экспорта (сообщение по-русски)."""


@dataclass
class _Turn:
    speaker: str | None
    start: float
    texts: list[str]


# --- время -------------------------------------------------------------------------


def _hms(seconds: float) -> str:
    total = max(0, int(seconds))  # отбрасываем доли, как плееры
    return f"{total // 3600:02d}:{total % 3600 // 60:02d}:{total % 60:02d}"


def _duration_label(seconds: float) -> str:
    total = max(0, int(round(seconds)))
    h, m, s = total // 3600, total % 3600 // 60, total % 60
    return f"{h}:{m:02d}:{s:02d}" if h else f"{m}:{s:02d}"


def _cue_time(seconds: float, sep: str) -> str:
    ms_total = max(0, int(round(seconds * 1000)))
    h, rem = divmod(ms_total, 3_600_000)
    m, rem = divmod(rem, 60_000)
    s, ms = divmod(rem, 1000)
    return f"{h:02d}:{m:02d}:{s:02d}{sep}{ms:03d}"


# --- общие помощники ---------------------------------------------------------------


def _names(transcript: dict[str, Any]) -> dict[str, str]:
    return {s["id"]: s.get("name") or s["id"] for s in transcript.get("speakers", [])}


def _speaker_name(names: dict[str, str], speaker: str | None) -> str | None:
    if speaker is None:
        return None
    return names.get(speaker, speaker)


def _segments(transcript: dict[str, Any]) -> list[dict[str, Any]]:
    return [s for s in transcript.get("segments", []) if (s.get("text") or "").strip()]


def _turns(transcript: dict[str, Any]) -> list[_Turn]:
    """Подряд идущие сегменты одного спикера → одна реплика. Без спикеров — сегмент = абзац."""
    names = _names(transcript)
    turns: list[_Turn] = []
    for seg in _segments(transcript):
        name = _speaker_name(names, seg.get("speaker"))
        text = seg["text"].strip()
        if turns and name is not None and turns[-1].speaker == name:
            turns[-1].texts.append(text)
        else:
            turns.append(_Turn(name, float(seg.get("start", 0.0)), [text]))
    return turns


def _title(transcript: dict[str, Any]) -> str:
    return transcript.get("file_name") or "Транскрипт"


def _meta_line(transcript: dict[str, Any]) -> str:
    created = transcript.get("created_at") or time.time()
    parts = [
        f"Дата: {time.strftime('%d.%m.%Y %H:%M', time.localtime(created))}",
        f"Длительность: {_duration_label(float(transcript.get('duration') or 0.0))}",
    ]
    speakers = [s.get("name") or s["id"] for s in transcript.get("speakers", [])]
    if speakers:
        parts.append(f"Спикеры: {', '.join(speakers)}")
    return " · ".join(parts)


def default_filename(transcript: dict[str, Any], fmt: str) -> str:
    stem = Path(transcript.get("file_name") or "transcript").stem or "transcript"
    return f"{stem}.{fmt}"


# --- форматы -----------------------------------------------------------------------


def render_txt(transcript: dict[str, Any], include_timestamps: bool = True) -> str:
    """``[00:01:23] Иван: текст``; пустая строка при смене спикера."""
    names = _names(transcript)
    lines: list[str] = []
    prev: object = object()
    for seg in _segments(transcript):
        name = _speaker_name(names, seg.get("speaker"))
        if lines and name != prev:
            lines.append("")
        prev = name
        prefix = f"[{_hms(seg.get('start', 0.0))}] " if include_timestamps else ""
        speaker = f"{name}: " if name else ""
        lines.append(f"{prefix}{speaker}{seg['text'].strip()}")
    return "\n".join(lines) + ("\n" if lines else "")


def render_md(transcript: dict[str, Any], include_timestamps: bool = True) -> str:
    out = [f"# {_title(transcript)}", "", f"_{_meta_line(transcript)}_", ""]
    for turn in _turns(transcript):
        stamp = f"[{_hms(turn.start)}]" if include_timestamps else ""
        if turn.speaker:
            out.append(f"**{turn.speaker}** {stamp}".rstrip())
            out.append(" ".join(turn.texts))
        else:
            out.append(f"{stamp} {' '.join(turn.texts)}".strip())
        out.append("")
    return "\n".join(out)


def _render_cues(transcript: dict[str, Any], sep: str, numbered: bool) -> list[str]:
    names = _names(transcript)
    out: list[str] = []
    for i, seg in enumerate(_segments(transcript), start=1):
        name = _speaker_name(names, seg.get("speaker"))
        text = seg["text"].strip()
        if numbered:
            out.append(str(i))
        out.append(f"{_cue_time(seg.get('start', 0.0), sep)} --> {_cue_time(seg.get('end', 0.0), sep)}")
        out.append(f"{name}: {text}" if name else text)
        out.append("")
    return out


def render_srt(transcript: dict[str, Any]) -> str:
    return "\n".join(_render_cues(transcript, ",", numbered=True))


def render_vtt(transcript: dict[str, Any]) -> str:
    return "\n".join(["WEBVTT", "", *_render_cues(transcript, ".", numbered=False)])


def render_json(transcript: dict[str, Any]) -> str:
    names = _names(transcript)
    data = {k: v for k, v in transcript.items() if k not in ("media_url", "segments")}
    data["segments"] = [
        {**seg, "speaker_id": seg.get("speaker"), "speaker": _speaker_name(names, seg.get("speaker"))}
        for seg in transcript.get("segments", [])
    ]
    data["full_text"] = " ".join(s["text"].strip() for s in _segments(transcript))
    return json.dumps(data, ensure_ascii=False, indent=2)


def _set_run_font(run, name: str, fallback: str) -> None:
    from docx.oxml.ns import qn

    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rfonts.set(qn(attr), name)
    rfonts.set(qn("w:cs"), fallback)


def _register_font_fallbacks(document, fallbacks: dict[str, str]) -> None:
    """Добавляет ``<w:altName>`` в fontTable: Word подставит запасной шрифт, если основного нет."""
    from lxml import etree

    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    for rel in document.part.rels.values():
        if rel.is_external or not rel.reltype.endswith("/fontTable"):
            continue
        part = rel.target_part
        try:
            root = etree.fromstring(part.blob)
            existing = {f.get(f"{{{w}}}name") for f in root.findall(f"{{{w}}}font")}
            for name, alt in fallbacks.items():
                if name in existing:
                    continue
                font = etree.SubElement(root, f"{{{w}}}font", {f"{{{w}}}name": name})
                etree.SubElement(font, f"{{{w}}}altName", {f"{{{w}}}val": alt})
                etree.SubElement(font, f"{{{w}}}family", {f"{{{w}}}val": "swiss"})
                etree.SubElement(font, f"{{{w}}}pitch", {f"{{{w}}}val": "variable"})
            part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        except Exception:  # noqa: BLE001 — запасные шрифты не критичны
            pass
        return


def _write_docx(transcript: dict[str, Any], path: Path, include_timestamps: bool) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    doc = Document()
    _register_font_fallbacks(doc, {"Montserrat": "Arial", "Calibri": "Arial"})

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rfonts.set(qn(attr), "Calibri")
    rfonts.set(qn("w:cs"), "Arial")

    title = doc.add_paragraph()
    title_run = title.add_run(_title(transcript))
    _set_run_font(title_run, "Montserrat", "Arial")
    title_run.bold = True
    title_run.font.size = Pt(18)
    title_run.font.color.rgb = RGBColor.from_string(NAVY)
    title.paragraph_format.space_after = Pt(4)

    meta = doc.add_paragraph()
    meta_run = meta.add_run(_meta_line(transcript))
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = RGBColor.from_string(GRAY)
    meta.paragraph_format.space_after = Pt(14)

    for turn in _turns(transcript):
        if turn.speaker or include_timestamps:
            head = doc.add_paragraph()
            head.paragraph_format.space_before = Pt(8)
            head.paragraph_format.space_after = Pt(2)
            head.paragraph_format.keep_with_next = True
            if turn.speaker:
                name_run = head.add_run(turn.speaker)
                name_run.bold = True
                name_run.font.color.rgb = RGBColor.from_string(STEEL_BLUE)
            if include_timestamps:
                stamp = head.add_run(("  " if turn.speaker else "") + _hms(turn.start))
                stamp.font.size = Pt(8.5)
                stamp.font.color.rgb = RGBColor.from_string(GRAY)
        body = doc.add_paragraph(" ".join(turn.texts))
        body.paragraph_format.space_after = Pt(4)

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run(FOOTER_TEXT)
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor.from_string(GRAY)

    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(f".{path.name}.tmp")
    try:
        doc.save(str(tmp))
        tmp.replace(path)
    finally:
        if tmp.exists():
            tmp.unlink()


def render(transcript: dict[str, Any], fmt: str, include_timestamps: bool = True) -> str:
    """Текстовые форматы (всё, кроме docx)."""
    if fmt == "txt":
        return render_txt(transcript, include_timestamps)
    if fmt == "md":
        return render_md(transcript, include_timestamps)
    if fmt == "srt":
        return render_srt(transcript)
    if fmt == "vtt":
        return render_vtt(transcript)
    if fmt == "json":
        return render_json(transcript)
    raise ExportError(f"Неизвестный формат экспорта: {fmt}")


def export(
    transcript: dict[str, Any], fmt: str, path: str | Path, include_timestamps: bool = True
) -> Path:
    """Пишет файл ``path`` в формате ``fmt``; возвращает путь."""
    if fmt not in FORMATS:
        raise ExportError(f"Неизвестный формат экспорта: {fmt}")
    path = Path(path)
    try:
        if fmt == "docx":
            _write_docx(transcript, path, include_timestamps)
        else:
            atomic_write_text(path, render(transcript, fmt, include_timestamps))
    except PermissionError as exc:
        raise ExportError(f"Нет доступа для записи: {path}") from exc
    except OSError as exc:
        raise ExportError(f"Не удалось сохранить файл {path.name}: {exc.strerror or exc}") from exc
    return path


def unique_path(directory: Path, filename: str, taken: Iterable[Path] = ()) -> Path:
    """``name.ext`` → ``name (2).ext`` … если файл уже существует."""
    taken_set = {p.resolve() for p in taken}
    candidate = directory / filename
    stem, suffix = Path(filename).stem, Path(filename).suffix
    n = 2
    while candidate.exists() or candidate.resolve() in taken_set:
        candidate = directory / f"{stem} ({n}){suffix}"
        n += 1
    return candidate
