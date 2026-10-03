from __future__ import annotations

import json
import re

import pytest

from bp_transcriber import exporters
from bp_transcriber.history import HistoryStore


@pytest.fixture
def transcript(tmp_path, result):
    store = HistoryStore(tmp_path / "h")
    tid = store.save_result(result, "/data/Планёрка 03.10.mp3")
    return store.rename_speaker(tid, "S1", "Иван")


def test_txt(transcript):
    txt = exporters.render_txt(transcript)
    lines = txt.splitlines()
    assert lines[0] == "[00:00:00] Иван: Добрый день, коллеги."
    assert lines[1] == "[00:00:02] Иван: Начнём с итогов квартала."
    assert lines[2] == ""  # смена спикера
    assert lines[3] == "[00:00:04] Спикер 2: Продажи выросли на ёлочные игрушки."
    assert lines[5] == "[01:01:01] Иван: Отлично, спасибо."
    plain = exporters.render_txt(transcript, include_timestamps=False)
    assert plain.splitlines()[0] == "Иван: Добрый день, коллеги."
    assert "[" not in plain


def test_md(transcript):
    md = exporters.render_md(transcript)
    assert md.startswith("# Планёрка 03.10.mp3\n")
    assert "Длительность: 1:01:03" in md and "Спикеры: Иван, Спикер 2" in md
    assert "**Иван** [00:00:00]\nДобрый день, коллеги. Начнём с итогов квартала." in md
    assert "**Спикер 2** [00:00:04]" in md


def test_srt_and_vtt(transcript):
    srt = exporters.render_srt(transcript)
    blocks = srt.strip().split("\n\n")
    assert blocks[0].splitlines() == ["1", "00:00:00,500 --> 00:00:02,000", "Иван: Добрый день, коллеги."]
    assert "01:01:01,001 --> 01:01:02,500" in srt  # без ошибки округления float
    for block in blocks:
        assert re.match(r"\d+\n\d{2}:\d{2}:\d{2},\d{3} --> \d{2}:\d{2}:\d{2},\d{3}\n.+", block)
    vtt = exporters.render_vtt(transcript)
    assert vtt.startswith("WEBVTT\n\n00:00:00.500 --> 00:00:02.000\nИван: Добрый день, коллеги.")


def test_json(transcript):
    data = json.loads(exporters.render_json(transcript))
    assert data["segments"][0]["speaker"] == "Иван"
    assert data["segments"][0]["speaker_id"] == "S1"
    assert data["speakers"][0]["name"] == "Иван"
    assert "media_url" not in data
    assert data["full_text"].startswith("Добрый день")


def test_docx(transcript, tmp_path):
    from docx import Document

    path = exporters.export(transcript, "docx", tmp_path / "out" / "t.docx")
    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Планёрка 03.10.mp3" in doc.paragraphs[0].text
    assert "Иван" in text and "Спикер 2" in text
    assert "Продажи выросли на ёлочные игрушки." in text
    title_run = doc.paragraphs[0].runs[0]
    assert title_run.bold and str(title_run.font.color.rgb) == "0B1D3A"
    assert title_run.font.name == "Montserrat"
    name_runs = [r for p in doc.paragraphs for r in p.runs if r.text == "Иван"]
    assert name_runs and str(name_runs[0].font.color.rgb) == "1E3A5F" and name_runs[0].bold
    footer = doc.sections[0].footer.paragraphs[0].text
    assert footer == "Создано в BP Transcriber · bestpracticeai.ru"
    assert not any(p.name.endswith(".tmp") for p in path.parent.iterdir())


@pytest.mark.parametrize("fmt", ["txt", "md", "docx", "srt", "vtt", "json"])
def test_export_all_formats(transcript, tmp_path, fmt):
    name = exporters.default_filename(transcript, fmt)
    assert name == f"Планёрка 03.10.{fmt}"
    path = exporters.export(transcript, fmt, tmp_path / name)
    assert path.is_file() and path.stat().st_size > 0


def test_unknown_format(transcript, tmp_path):
    with pytest.raises(exporters.ExportError):
        exporters.export(transcript, "pdf", tmp_path / "x.pdf")


def test_unique_path(tmp_path):
    (tmp_path / "a.txt").write_text("x")
    (tmp_path / "a (2).txt").write_text("x")
    assert exporters.unique_path(tmp_path, "a.txt").name == "a (3).txt"
    assert exporters.unique_path(tmp_path, "b.txt").name == "b.txt"


def test_no_speakers(tmp_path, result):
    for seg in result.segments:
        seg.speaker = None
    store = HistoryStore(tmp_path / "h2")
    t = store.load(store.save_result(result, "/x/a.wav"))
    assert t["speakers"] == []
    txt = exporters.render_txt(t)
    assert txt.splitlines()[0] == "[00:00:00] Добрый день, коллеги."
    assert "" not in txt.splitlines()
    assert exporters.render_srt(t).splitlines()[2] == "Добрый день, коллеги."
